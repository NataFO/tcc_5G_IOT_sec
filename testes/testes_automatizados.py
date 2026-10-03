"""
Testes automatizados da API (FastAPI + LSTM) — CT02, CT03, CT04, CT11, CT12
(+ uma verificação de autenticação que serve de evidência parcial do CT13).

O que ele faz:
  1. Confere que POST /logs/ SEM token devolve 401 (parte do CT13).
  2. Faz login (POST /auth/login) e pega o token JWT.
  3. Garante um dispositivo dedicado "Sensor-Testes-Automatizados" (cria se
     não existir), para os logs/alertas do teste ficarem separados dos demais.
  4. Faz 3 requisições de aquecimento (descartadas) e envia as amostras de testes/conjunto_teste_bruto.csv, UMA A UMA, para
     POST /logs/ — com os valores BRUTOS das 20 features, como faria uma
     ferramenta de captura real — e mede o tempo de cada requisição.
  5. Compara a classificação da API com o rótulo real e calcula:
       CT02  acurácia                     (critério: >= 95%)
       CT03  tempo de inferência P99      (critério: <= 100 ms)
       CT04  alerta gerado p/ cada ataque + tempo P95 até o alerta ser
             gravado no banco             (critério: <= 200 ms)
       CT11  FPR (falsos positivos)       (critério: < 2%)
       CT12  FNR em ataques DDoS          (critério: < 5%)
  6. Salva os resultados em testes/resultados/ e, com --matriz, preenche as
     linhas desses CTs na matriz_testes_CT01-CT14.xlsx.

Sobre a latência: o backend devolve, no cabeçalho padrão "Server-Timing",
quanto tempo cada etapa levou DENTRO do servidor (consulta ao banco,
inferência do LSTM, gravação do log, gravação do alerta). O script usa:
  - CT03: só a etapa de inferência (é o que o critério pede);
  - CT04: o tempo no servidor desde a chegada da requisição até o alerta
    estar gravado (só nas amostras classificadas como ataque).
Ele também mede a requisição inteira do lado do cliente (inclui rede) e
mostra a diferença, para ficar claro quanto é rede e quanto é servidor.
Se o backend não mandar Server-Timing (versão antiga), o script cai para a
medida ponta a ponta e avisa.

ATENÇÃO: cada amostra enviada vira uma linha em tb_log_rede (e, se for
classificada como ataque, uma em tb_alerta), todas no dispositivo
"Sensor-Testes-Automatizados". Rode primeiro com poucas amostras.

Uso (dentro da pasta do projeto, com o backend rodando):
  pip install requests pandas openpyxl
  python testes/testes_automatizados.py --amostras 50                 # teste rápido local
  python testes/testes_automatizados.py --amostras todas --matriz matriz_testes_CT01-CT14.xlsx
  python testes/testes_automatizados.py --url https://SEU-BACKEND.up.railway.app --amostras 1000

A senha é pedida no terminal (não fica salva em lugar nenhum). Também pode
vir da variável de ambiente TCC_SENHA.
"""
import argparse
import getpass
import json
import os
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

try:
    import requests
except ImportError:
    sys.exit("Falta a biblioteca 'requests'. Rode: pip install requests")

PASTA = Path(__file__).resolve().parent
CSV_PADRAO = PASTA / "conjunto_teste_bruto.csv"
NOME_DISPOSITIVO = "Sensor-Testes-Automatizados"

# Critérios de aceite (Tabela 11 do TCC)
CRITERIO_ACURACIA = 0.95   # CT02
CRITERIO_P99_MS = 100.0    # CT03
CRITERIO_ALERTA_MS = 200.0 # CT04
CRITERIO_FPR = 0.02        # CT11
CRITERIO_FNR = 0.05        # CT12


# ── helpers ───────────────────────────────────────────────────────────────
def pct(x):
    return f"{x * 100:.2f}%".replace(".", ",")


def ms(x):
    return f"{x:.1f} ms".replace(".", ",")


def falhar(msg, resp=None):
    if resp is not None:
        msg += f"\n  HTTP {resp.status_code}: {resp.text[:300]}"
    sys.exit("ERRO: " + msg)


def montar_payload(linha: pd.Series, i: int, id_dispositivo: int, features: list) -> dict:
    """Monta o corpo do POST /logs/: metadados do fluxo + as 20 features brutas."""
    pacotes = int(max(linha["network_packets_all_count"], 0))
    tamanho = float(max(linha["network_packet_size_max"], 0))
    payload = {
        "id_dispositivo": id_dispositivo,
        "ip_origem": f"10.99.{(i // 250) % 250}.{i % 250 + 1}",
        "ip_destino": "10.0.0.1",
        "porta_origem": 40000 + (i % 20000),
        "porta_destino": 80,
        "protocolo": "TCP",
        "pacotes_enviados": pacotes,
        "bytes_enviados": int(pacotes * tamanho),
        "duracao_fluxo_ms": 5000.0,  # cada amostra do dataset é uma janela de 5 s
    }
    for f in features:
        payload[f] = float(linha[f])
    return payload


def status_ct(ok: bool) -> str:
    return "Aprovado" if ok else "Reprovado"


def ler_server_timing(cabecalho: str) -> dict:
    """'inferencia;dur=12.3, total;dur=40.1' -> {'inferencia': 12.3, 'total': 40.1}"""
    tempos = {}
    for parte in (cabecalho or "").split(","):
        nome, _, resto = parte.strip().partition(";")
        if resto.startswith("dur="):
            try:
                tempos[nome] = float(resto[4:])
            except ValueError:
                pass
    return tempos


def percentis(valores) -> tuple:
    return tuple(float(x) for x in np.percentile(valores, [50, 95, 99]))


# ── etapas ────────────────────────────────────────────────────────────────
def checar_sem_token(s: requests.Session, url: str) -> int:
    r = s.post(f"{url}/logs/", json={}, timeout=30)
    return r.status_code


def fazer_login(s: requests.Session, url: str, login: str, senha: str) -> str:
    r = s.post(f"{url}/auth/login", json={"login": login, "senha": senha}, timeout=30)
    if r.status_code != 200:
        falhar("login falhou — confira usuário/senha.", r)
    return r.json()["access_token"]


def obter_dispositivo(s: requests.Session, url: str) -> int:
    r = s.get(f"{url}/dispositivos/", timeout=30)
    if r.status_code != 200:
        falhar("não consegui listar dispositivos.", r)
    for d in r.json():
        if d["nome_dispositivo"] == NOME_DISPOSITIVO:
            return d["id_dispositivo"]
    r = s.post(f"{url}/dispositivos/", json={
        "nome_dispositivo": NOME_DISPOSITIVO,
        "tipo": "sensor",
        "ip_address": "10.99.99.99",
        "gnodeb_associado": "gNodeB-Testes",
    }, timeout=30)
    if r.status_code != 200:
        falhar(f"não consegui criar o dispositivo '{NOME_DISPOSITIVO}'.", r)
    print(f"  Dispositivo '{NOME_DISPOSITIVO}' criado (id {r.json()['id_dispositivo']}).")
    return r.json()["id_dispositivo"]


def contar_alertas(s: requests.Session, url: str, id_dispositivo: int, desde: str) -> int:
    """Alertas do dispositivo de teste com data_hora_alerta >= 'desde' (relógio do servidor)."""
    r = s.get(f"{url}/alertas/", params={"limite": 1_000_000}, timeout=120)
    if r.status_code != 200:
        falhar("não consegui listar alertas.", r)
    return sum(
        1 for a in r.json()
        if a["id_dispositivo"] == id_dispositivo and a["data_hora_alerta"] >= desde
    )


def preencher_matriz(caminho: Path, linhas: dict):
    try:
        import openpyxl
    except ImportError:
        print("  (openpyxl não instalado — pulei a matriz. pip install openpyxl)")
        return
    backup = caminho.with_suffix(".backup.xlsx")
    shutil.copy(caminho, backup)
    wb = openpyxl.load_workbook(caminho)
    ws = wb["CT01-CT14"]
    for row in range(8, 22):
        ct = ws.cell(row, 1).value
        if ct in linhas:
            for col, valor in zip("FGHIJK", linhas[ct]):
                ws[f"{col}{row}"] = valor
    try:
        wb.save(caminho)
    except PermissionError:
        sys.exit(f"ERRO: não consegui salvar {caminho.name} — feche o arquivo no Excel e rode de novo.")
    print(f"  Matriz atualizada: {caminho} (cópia anterior em {backup.name})")


# ── principal ─────────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser(description="Testes automatizados CT02/CT03/CT04/CT11/CT12")
    ap.add_argument("--url", default="http://localhost:8000", help="URL do backend (padrão: local)")
    ap.add_argument("--login", default="admin")
    ap.add_argument("--senha", default=None, help="se omitida, é pedida no terminal")
    ap.add_argument("--amostras", default="200",
                    help="quantas amostras enviar (número) ou 'todas' (6.759). Padrão: 200")
    ap.add_argument("--csv", default=str(CSV_PADRAO))
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--matriz", default=None,
                    help="caminho da matriz_testes_CT01-CT14.xlsx para preencher CT02/03/04/11/12")
    args = ap.parse_args()

    url = args.url.rstrip("/")
    ambiente = "Local" if any(h in url for h in ("localhost", "127.0.0.1")) else "Nuvem (Railway)"

    df = pd.read_csv(args.csv)
    features = [c for c in df.columns if c.startswith("network_")]
    if len(features) != 20:
        falhar(f"esperava 20 features no CSV, achei {len(features)}.")
    if args.amostras != "todas":
        n = int(args.amostras)
        if n < len(df):
            # amostra estratificada: mantém a proporção benigno/ataque do conjunto de teste
            partes = [g.sample(max(1, round(n * len(g) / len(df))), random_state=args.seed)
                      for _, g in df.groupby("label")]
            df = pd.concat(partes).sample(frac=1, random_state=args.seed).reset_index(drop=True)

    senha = args.senha or os.getenv("TCC_SENHA") or getpass.getpass(f"Senha do usuário '{args.login}': ")

    print(f"\nBackend: {url}  ({ambiente})")
    s = requests.Session()
    try:
        r = s.get(url + "/", timeout=60)
        r.raise_for_status()
    except Exception as e:
        falhar(f"backend não respondeu em {url}/ — ele está rodando? ({e})")

    print("[1/4] Verificando que a API recusa requisição sem token...")
    cod_sem_token = checar_sem_token(s, url)
    print(f"  POST /logs/ sem token -> HTTP {cod_sem_token} "
          f"({'ok' if cod_sem_token == 401 else 'ESPERAVA 401'})")

    print("[2/4] Login e dispositivo de teste...")
    token = fazer_login(s, url, args.login, senha)
    s.headers["Authorization"] = f"Bearer {token}"
    id_disp = obter_dispositivo(s, url)

    total = len(df)
    print(f"[3/4] Enviando {total} amostras (benigno: {(df.label == 0).sum()}, "
          f"ataque: {(df.label == 1).sum()}) para POST /logs/ ...")
    # Aquecimento: as primeiras chamadas ao TensorFlow são bem mais lentas
    # (carregamento/compilação do grafo). 3 requisições descartadas antes de medir.
    for i in range(3):
        r = s.post(f"{url}/logs/", json=montar_payload(df.iloc[i], i, id_disp, features), timeout=120)
        if r.status_code != 200:
            falhar("POST /logs/ falhou no aquecimento.", r)
    inicio = datetime.now()
    resultados = []
    primeiro_momento = None
    for i, linha in df.iterrows():
        payload = montar_payload(linha, i, id_disp, features)
        t0 = time.perf_counter()
        r = s.post(f"{url}/logs/", json=payload, timeout=60)
        lat = (time.perf_counter() - t0) * 1000
        if r.status_code != 200:
            falhar(f"POST /logs/ falhou na amostra {i}.", r)
        resp = r.json()
        tempos = ler_server_timing(r.headers.get("Server-Timing"))
        if primeiro_momento is None:
            primeiro_momento = resp["momento_captura"]
        resultados.append({
            "amostra": i,
            "label_real": int(linha["label"]),
            "tipo_ataque_real": linha["tipo_ataque_real"],
            "classificacao_api": resp["classificacao_ia"],
            "probabilidade_api": resp["probabilidade_ia"],
            "predito": 0 if resp["classificacao_ia"] == "Normal" else 1,
            "latencia_ms": round(lat, 2),
            "t_consulta_ms": tempos.get("consulta"),
            "t_inferencia_ms": tempos.get("inferencia"),
            "t_modelo_ms": tempos.get("modelo"),
            "t_fila_ms": tempos.get("fila_modelo"),
            "t_gravar_log_ms": tempos.get("gravar_log"),
            "t_gravar_alerta_ms": tempos.get("gravar_alerta"),
            "t_servidor_ms": tempos.get("total"),
            "id_log": resp["id_log"],
        })
        if (len(resultados)) % 50 == 0 or len(resultados) == total:
            print(f"  {len(resultados)}/{total}", end="\r", flush=True)
    print()
    fim = datetime.now()
    res = pd.DataFrame(resultados)

    print("[4/4] Conferindo alertas gerados...")
    alertas = contar_alertas(s, url, id_disp, primeiro_momento)

    # ── métricas ──
    y, p = res.label_real.values, res.predito.values
    tp = int(((p == 1) & (y == 1)).sum()); tn = int(((p == 0) & (y == 0)).sum())
    fp = int(((p == 1) & (y == 0)).sum()); fn = int(((p == 0) & (y == 1)).sum())
    acuracia = (tp + tn) / len(res)
    fpr = fp / (fp + tn) if (fp + tn) else float("nan")
    fnr = fn / (fn + tp) if (fn + tp) else float("nan")
    lat = res.latencia_ms.values
    p50, p95, p99 = np.percentile(lat, [50, 95, 99])  # requisição inteira (cliente)
    tem_timing = res.t_inferencia_ms.notna().all()
    if tem_timing:
        inf50, inf95, inf99 = percentis(res.t_inferencia_ms)
        fonte_ct03 = "tempo de inferência medido no servidor"
    else:
        inf50, inf95, inf99 = p50, p95, p99
        fonte_ct03 = "requisição completa (backend sem Server-Timing)"
    ddos = res[res.tipo_ataque_real == "ddos"]
    fnr_ddos = (ddos.predito == 0).mean() if len(ddos) else float("nan")
    ataques = res[res.label_real == 1].assign(perdido=lambda d: d.predito == 0)
    fnr_por_tipo = ataques.groupby("tipo_ataque_real")["perdido"].agg(n="size", fnr="mean")
    col_alerta = "t_servidor_ms" if tem_timing else "latencia_ms"
    lat_ataques = res[res.predito == 1][col_alerta].values
    p95_alerta = float(np.percentile(lat_ataques, 95)) if len(lat_ataques) else float("nan")
    fonte_ct04 = ("tempo no servidor até o alerta gravado" if tem_timing
                  else "requisição completa (backend sem Server-Timing)")
    previstos_ataque = int((p == 1).sum())

    # ── relatório no terminal ──
    print("\n" + "=" * 66)
    print(f"RESULTADOS — {len(res)} amostras — {ambiente} — {inicio:%d/%m/%Y %H:%M}")
    print("=" * 66)
    print(f"Matriz de confusão:  VP={tp}  VN={tn}  FP={fp}  FN={fn}")
    print(f"CT02 Acurácia:         {pct(acuracia):>10}   (critério >= 95%)   -> {status_ct(acuracia >= CRITERIO_ACURACIA)}")
    print(f"CT03 Inferência P99:   {ms(inf99):>10}   (critério <= 100 ms) -> {status_ct(inf99 <= CRITERIO_P99_MS)}"
          f"   [P50 {ms(inf50)} | P95 {ms(inf95)}]")
    alerta_ok = alertas == previstos_ataque and p95_alerta <= CRITERIO_ALERTA_MS
    print(f"CT04 Alertas:          {alertas}/{previstos_ataque} gerados; P95 até gravar {ms(p95_alerta)}"
          f"   (critério <= 200 ms) -> {status_ct(alerta_ok)}")
    print(f"CT11 FPR:              {pct(fpr):>10}   (critério < 2%)      -> {status_ct(fpr < CRITERIO_FPR)}")
    print(f"CT12 FNR (DDoS, n={len(ddos)}): {pct(fnr_ddos):>7}   (critério < 5%)      -> {status_ct(fnr_ddos < CRITERIO_FNR)}")
    print(f"     FNR geral (todos os ataques): {pct(fnr)}")
    print("     FNR por tipo de ataque:")
    for tipo, row in fnr_por_tipo.sort_values("n", ascending=False).iterrows():
        print(f"       {tipo:<12} n={int(row.n):<5} FNR={pct(row.fnr)}")
    print(f"Autenticação: POST /logs/ sem token -> HTTP {cod_sem_token}")
    print("\nOnde o tempo foi gasto (ms)          P50        P95        P99")
    etapas = [("Requisição inteira (cliente)", res.latencia_ms)]
    if tem_timing:
        etapas += [
            ("  consulta ao banco", res.t_consulta_ms),
            ("  inferência LSTM", res.t_inferencia_ms),
            ("    (execução do modelo)", res.t_modelo_ms.dropna()),
            ("    (espera na fila)", res.t_fila_ms.dropna()),
            ("  gravar log", res.t_gravar_log_ms),
            ("  gravar alerta (só ataques)", res[res.predito == 1].t_gravar_alerta_ms),
            ("  total no servidor", res.t_servidor_ms),
            ("  rede + HTTP (diferença)", res.latencia_ms - res.t_servidor_ms),
        ]
    else:
        print("  (backend sem Server-Timing: só a medida ponta a ponta disponível)")
    for nome, serie in etapas:
        if len(serie):
            a, b, c = percentis(serie)
            print(f"  {nome:<32} {ms(a):>10} {ms(b):>10} {ms(c):>10}")
    if len(res) < 1000:
        print("\nObs.: poucas amostras — para o TCC, rode com --amostras todas (métricas mais estáveis).")

    # ── arquivos ──
    pasta_res = PASTA / "resultados"
    pasta_res.mkdir(exist_ok=True)
    tag = f"{inicio:%Y%m%d_%H%M%S}_{'local' if ambiente == 'Local' else 'nuvem'}"
    arq_det = pasta_res / f"detalhado_{tag}.csv"
    arq_res = pasta_res / f"resumo_{tag}.json"
    res.to_csv(arq_det, index=False)
    resumo = {
        "backend": url, "ambiente": ambiente,
        "inicio": inicio.isoformat(timespec="seconds"), "fim": fim.isoformat(timespec="seconds"),
        "amostras": len(res), "vp": tp, "vn": tn, "fp": fp, "fn": fn,
        "acuracia": acuracia, "fpr": fpr, "fnr_geral": fnr, "fnr_ddos": fnr_ddos, "n_ddos": len(ddos),
        "fnr_por_tipo": {k: {"n": int(v.n), "fnr": float(v.fnr)} for k, v in fnr_por_tipo.iterrows()},
        "latencia_requisicao_ms": {"p50": p50, "p95": p95, "p99": p99, "media": float(lat.mean()), "max": float(lat.max())},
        "inferencia_ms": {"p50": inf50, "p95": inf95, "p99": inf99, "fonte": fonte_ct03},
        "alertas_gerados": alertas, "previstos_ataque": previstos_ataque,
        "p95_ate_alerta_ms": p95_alerta, "fonte_ct04": fonte_ct04,
        "http_sem_token": cod_sem_token,
    }
    arq_res.write_text(json.dumps(resumo, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nSalvo: {arq_det}\nSalvo: {arq_res}")

    if args.matriz:
        quando = f"{inicio:%d/%m/%Y %H:%M}"
        evid = f"testes/resultados/{arq_res.name}"
        obs = f"{len(res)} amostras do conjunto de teste (valores brutos) via POST /logs/; testes_automatizados.py"
        linhas = {
            "CT02": (ambiente, quando, f"Acurácia {pct(acuracia)} (VP={tp}, VN={tn}, FP={fp}, FN={fn})",
                     status_ct(acuracia >= CRITERIO_ACURACIA), evid, obs),
            "CT03": (ambiente, quando, f"Inferência P99 {ms(inf99)} (P50 {ms(inf50)}, P95 {ms(inf95)})",
                     status_ct(inf99 <= CRITERIO_P99_MS), evid,
                     f"{obs}. Fonte: {fonte_ct03}; requisição inteira P99 {ms(p99)}"),
            "CT04": (ambiente, quando, f"{alertas}/{previstos_ataque} alertas gerados; P95 até gravar {ms(p95_alerta)}",
                     status_ct(alerta_ok), evid, f"{obs}. Fonte: {fonte_ct04}"),
            "CT11": (ambiente, quando, f"FPR {pct(fpr)} ({fp} de {fp + tn} benignos)",
                     status_ct(fpr < CRITERIO_FPR), evid, obs),
            "CT12": (ambiente, quando, f"FNR DDoS {pct(fnr_ddos)} (n={len(ddos)}); FNR geral {pct(fnr)}",
                     status_ct(fnr_ddos < CRITERIO_FNR), evid, obs),
        }
        preencher_matriz(Path(args.matriz), linhas)


if __name__ == "__main__":
    main()
